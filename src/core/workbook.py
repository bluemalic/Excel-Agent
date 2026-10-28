"""WorkbookManager：workbook 唯一句柄（单一真相源）。

职责（规格 §7 铁律①）：
- 所有写操作经本管理器串行执行（全局线程锁），版本链与逆操作日志严格有序；
- plan（生成 diff 预览）与 apply（确认后落盘）分离，人在回路；
- undo / redo / 跳转第 N 步 = 操作游标移动 + 从头重放 sheet 版本指针与样式表；
- 导出 xlsx：v0 且源自上传/示例时直接复制原始文件字节（AC-5 逐字节一致），
  其余情况用 openpyxl 重建并套用样式表。
"""

from __future__ import annotations

import copy
import hashlib
import shutil
import threading
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from src import config
from src.core import mutations
from src.core.diff import DiffResult
from src.core.importer import import_file
from src.core.schema_extract import SheetSchema, extract_schema
from src.core.storage import Storage, get_storage


class WorkbookNotFoundError(KeyError):
    pass


@dataclass
class WriteOp:
    """一次结构化写操作计划（LLM 产出参数 → 交给本管理器校验执行）。"""
    op_type: str
    sheet: str
    params: dict = field(default_factory=dict)
    description: str = ""
    new_sheet_df: pd.DataFrame | None = None   # 仅 create_sheet：新表初始内容

    def describe(self) -> str:
        return self.description or f"{self.op_type} @ {self.sheet}"


@dataclass
class PlanResult:
    """写操作执行前的产物：预览 diff + 待执行结果。"""
    outcome: mutations.MutationOutcome
    diff: DiffResult
    op: WriteOp

    def forward_params(self) -> dict:
        """落盘时写入 ops 表的正向参数（样式操作展开为逐 ref 条目，供重放）。"""
        if self.op.op_type == "style_cells":
            style = self.op.params.get("style") or {}
            entries = {t: style for t in self.op.params.get("targets", [])}
            return {"sheet": self.op.sheet, "entries": entries,
                    "description": self.diff.description}
        params = dict(self.op.params)
        params["sheet"] = self.op.sheet
        params.setdefault("description", self.diff.description)
        return params


class WorkbookManager:
    _write_lock = threading.Lock()   # 铁律①：写操作全局串行
    _instance: "WorkbookManager | None" = None
    _instance_lock = threading.Lock()

    def __init__(self, storage: Storage | None = None):
        self.storage = storage or get_storage()
        config.ensure_dirs()
        self._df_cache: dict[tuple, pd.DataFrame] = {}

    @classmethod
    def instance(cls) -> "WorkbookManager":
        with cls._instance_lock:
            if cls._instance is None:
                cls._instance = cls()
            return cls._instance

    # ------------------------------------------------------------------
    # 创建 / 读取
    # ------------------------------------------------------------------
    def create_from_import(self, path: str | Path, display_name: str | None = None,
                           header_rows: dict[str, int] | None = None) -> tuple[str, list[str]]:
        """从文件创建 workbook，返回 (workbook_id, 警告列表)。"""
        result = import_file(path, sheet_header_rows=header_rows)
        wb_id = self.storage.create_workbook(
            name=display_name or Path(path).stem, origin="upload", original_path=str(path)
        )
        ws_dir = self._wb_dir(wb_id)
        for sheet_name, df in result.sheets.items():
            self._write_version(wb_id, sheet_name, 0, df, "导入初始版本")
        return wb_id, result.warnings

    def create_from_template(self, name: str, sheet_name: str,
                             columns: list[dict]) -> str:
        """模板建空表（C 档）：columns = [{name, type, ...}]。"""
        dtype_map = {
            "date": "datetime64[ns]", "datetime": "datetime64[ns]",
            "decimal": "Float64", "number": "Float64", "int": "Int64",
            "enum": "string", "string": "string", "text": "string", "bool": "boolean",
        }
        data = {}
        for col in columns:
            dt = dtype_map.get(str(col.get("type", "string")).lower(), "string")
            data[str(col["name"])] = pd.array([], dtype=dt)
        df = pd.DataFrame(data)
        wb_id = self.storage.create_workbook(name=name, origin="template", original_path=None)
        self._write_version(wb_id, sheet_name, 0, df, "模板初始空表")
        return wb_id

    def _wb_dir(self, wb_id: str) -> Path:
        d = config.WORKSPACE_DIR / wb_id
        d.mkdir(parents=True, exist_ok=True)
        return d

    def _sheet_file(self, wb_id: str, sheet: str) -> str:
        h = hashlib.sha1(sheet.encode("utf-8")).hexdigest()[:10]
        return str(self._wb_dir(wb_id) / f"sheet_{h}.parquet")

    def _write_version(self, wb_id: str, sheet: str, version: int, df: pd.DataFrame,
                       description: str, origin: str | None = None) -> str:
        stem = f"v{version}_" + hashlib.sha1(sheet.encode()).hexdigest()[:10] + ".parquet"
        path = str(self._wb_dir(wb_id) / stem)
        df.to_parquet(path, index=False)
        from src.core.schema_extract import extract_schema
        cols_json = _dumps([c.to_dict() for c in extract_schema(df, sheet).columns])
        self.storage.add_version(wb_id, sheet, version, path, cols_json, description)
        existing = self.storage.get_sheet(wb_id, sheet)
        sheet_origin = origin or (existing["origin"] if existing is not None else "import")
        self.storage.upsert_sheet(
            wb_id, sheet,
            n_rows=len(df), n_cols=len(df.columns),
            current_version=version, max_version=max(version, self._max_version(wb_id, sheet)),
            origin=sheet_origin,
        )
        self._df_cache.pop((wb_id, sheet, version), None)
        return path

    def _max_version(self, wb_id: str, sheet: str) -> int:
        s = self.storage.get_sheet(wb_id, sheet)
        return int(s["max_version"]) if s else 0

    def get_df(self, wb_id: str, sheet: str, version: int | None = None) -> pd.DataFrame:
        """读取 sheet 当前（或指定）版本的 DataFrame。"""
        s = self.storage.get_sheet(wb_id, sheet)
        if s is None:
            raise WorkbookNotFoundError(f"工作表「{sheet}」不存在")
        v = int(s["current_version"] if version is None else version)
        key = (wb_id, sheet, v)
        if key in self._df_cache:
            return self._df_cache[key]
        ver = self.storage.get_version(wb_id, sheet, v)
        if ver is None:
            raise WorkbookNotFoundError(f"工作表「{sheet}」版本 v{v} 缺失")
        df = pd.read_parquet(ver["parquet_path"])
        if len(self._df_cache) > 8:
            self._df_cache.clear()
        self._df_cache[key] = df
        return df

    def current_parquet_path(self, wb_id: str, sheet: str) -> str:
        s = self.storage.get_sheet(wb_id, sheet)
        if s is None:
            raise WorkbookNotFoundError(f"工作表「{sheet}」不存在")
        ver = self.storage.get_version(wb_id, sheet, int(s["current_version"]))
        return ver["parquet_path"]

    def get_schema(self, wb_id: str, sheet: str) -> SheetSchema:
        df = self.get_df(wb_id, sheet)
        return extract_schema(df, sheet)

    def list_sheet_names(self, wb_id: str) -> list[str]:
        return [r["sheet_name"] for r in self.storage.list_sheets(wb_id, only_active=True)]

    # ------------------------------------------------------------------
    # 写操作：plan（预览）与 apply（落盘）分离
    # ------------------------------------------------------------------
    def plan_write(self, wb_id: str, op: WriteOp) -> PlanResult:
        """执行写操作前生成 diff 预览（不改任何数据）。"""
        if op.op_type == "create_sheet":
            if not op.sheet or not op.sheet.strip():
                raise mutations.OpError("新工作表名称不能为空")
            new_df = op.new_sheet_df if op.new_sheet_df is not None else pd.DataFrame()
            outcome = mutations.EXECUTORS["create_sheet"](new_df, {"sheet_name": op.sheet,
                                                                   "description": op.description})
            return PlanResult(outcome=outcome, diff=outcome.diff, op=op)

        if op.op_type not in mutations.EXECUTORS:
            raise mutations.OpError(
                f"未知写操作类型：{op.op_type}。可用：{sorted(k for k in mutations.EXECUTORS if not k.startswith(('restore', 'style_restore', 'noop', 'drop_sheet')))}")
        params = dict(op.params)
        params["sheet"] = op.sheet
        params.setdefault("description", op.description)
        df = self.get_df(wb_id, op.sheet)
        # 列名先做一轮白名单校验（把正确列名回喂给模型）
        if op.op_type not in ("style_cells",):
            for key in ("column", "old"):
                if key in params and isinstance(params[key], str) and params[key] not in df.columns:
                    raise mutations.OpError(
                        f"列「{params[key]}」不存在。现有列：{list(df.columns)}")
        executor = mutations.EXECUTORS[op.op_type]
        outcome = executor(df, params)
        return PlanResult(outcome=outcome, diff=outcome.diff, op=op)

    def apply_write(self, wb_id: str, plan: PlanResult) -> dict:
        """用户确认后落盘：串行锁内执行，写新版本 + 记逆操作日志 + 前移游标。"""
        with WorkbookManager._write_lock:
            op = plan.op
            outcome = plan.outcome
            before = self._sheet_version(wb_id, op.sheet)

            # 若游标之后还有被撤销的操作，先截断 redo 分支（保留 1..cursor 的已应用前缀）
            self.storage.execute(
                "DELETE FROM ops WHERE workbook_id=? AND seq>?",
                (wb_id, self.storage.get_cursor(wb_id)),
            )

            if op.op_type == "create_sheet":
                existing = self.storage.get_sheet(wb_id, op.sheet)
                if existing is not None and existing["active"]:
                    raise mutations.OpError(f"工作表「{op.sheet}」已存在")
                if existing is None:
                    self._write_version(wb_id, op.sheet, 0, outcome.new_df,
                                        f"新建工作表：{plan.diff.description}", origin="op")
                    before = -1
                else:  # 曾被撤销停用：重新激活（数据取新建内容覆盖）
                    self._write_version(wb_id, op.sheet, self._max_version(wb_id, op.sheet) + 1,
                                        outcome.new_df, "重新创建工作表", origin="op")
                    before = int(existing["current_version"])
                after = self._sheet_version(wb_id, op.sheet)
            else:
                new_version = self._max_version(wb_id, op.sheet) + 1
                self._write_version(wb_id, op.sheet, new_version, outcome.new_df,
                                    plan.diff.description)
                after = new_version

            op_id = self.storage.add_op(
                wb_id=wb_id, sheet=op.sheet, op_type=op.op_type,
                params=plan.forward_params(),
                inverse=outcome.inverse,
                before_version=before, after_version=after,
                description=plan.diff.description,
            )
            self.storage.set_cursor(wb_id, self.storage.get_cursor(wb_id) + 1)
            # 统一重放：物化样式表、刷新行数/指针元数据（操作数少，代价可忽略）
            self._recompute(wb_id)
            return {"op_id": op_id, "version": after, "description": plan.diff.description,
                    "n_cells": plan.diff.n_cells}

    def _sheet_version(self, wb_id: str, sheet: str) -> int:
        s = self.storage.get_sheet(wb_id, sheet)
        return int(s["current_version"]) if s else -1

    # ------------------------------------------------------------------
    # 撤销 / 重做 / 跳转 / 快照（游标移动 + 重放）
    # ------------------------------------------------------------------
    def _ops(self, wb_id: str) -> list:
        return self.storage.list_ops(wb_id)

    def _recompute(self, wb_id: str) -> None:
        """按已应用前缀重放：sheet 版本指针 / 激活状态 / 样式表。"""
        cursor = self.storage.get_cursor(wb_id)
        ops = [o for o in self._ops(wb_id) if o["seq"] <= cursor]
        pointers: dict[str, int] = {}
        active: dict[str, bool] = {}
        styles: dict[str, dict[str, dict | None]] = {}
        for o in ops:
            sheet = o["sheet_name"]
            if sheet:
                pointers[sheet] = int(o["after_version"])
            if o["op_type"] == "create_sheet":
                active[sheet] = True
            elif o["op_type"] == "drop_sheet":
                active[sheet] = False
            elif o["op_type"] == "style_cells":
                entries = _loads(o["params_json"]).get("entries", {})
                styles.setdefault(sheet, {}).update(entries)
        for s in self.storage.list_sheets(wb_id, only_active=False):
            name = s["sheet_name"]
            if s["origin"] == "import":
                active.setdefault(name, True)
            if name not in pointers:
                pointers[name] = 0
        for name, v in pointers.items():
            s = self.storage.get_sheet(wb_id, name)
            if s is None:
                continue
            df_path_row = self.storage.get_version(wb_id, name, v)
            n_rows = n_cols = 0
            if df_path_row is not None:
                import pyarrow.parquet as pq
                meta = pq.read_metadata(df_path_row["parquet_path"])
                n_rows, n_cols = meta.num_rows, meta.num_columns
            self.storage.upsert_sheet(
                wb_id, name, n_rows=n_rows, n_cols=n_cols,
                current_version=v, max_version=int(s["max_version"]),
                active=1 if active.get(name, False) else 0,
                origin=s["origin"],
            )
        # 重建样式表（样式状态完全由已应用的 style_cells 决定）
        for s in self.storage.list_sheets(wb_id, only_active=False):
            self.storage.execute(
                "DELETE FROM styles WHERE workbook_id=? AND sheet_name=?",
                (wb_id, s["sheet_name"]))
        for sheet, entries in styles.items():
            for ref, style in entries.items():
                if style is not None:
                    self.storage.set_style(wb_id, sheet, ref, style)
        self._df_cache.clear()

    def history(self, wb_id: str) -> list[dict]:
        """操作历史（含游标位置，供 UI 时间线）。"""
        cursor = self.storage.get_cursor(wb_id)
        out = []
        for o in self._ops(wb_id):
            out.append({
                "seq": int(o["seq"]), "op_id": o["op_id"], "sheet": o["sheet_name"],
                "op_type": o["op_type"], "description": o["description"],
                "state": "已应用" if int(o["seq"]) <= cursor else "已撤销",
                "created_at": o["created_at"],
            })
        return out

    def undo(self, wb_id: str) -> str:
        with WorkbookManager._write_lock:
            cursor = self.storage.get_cursor(wb_id)
            if cursor <= 0:
                raise mutations.OpError("没有可撤销的操作")
            self.storage.set_cursor(wb_id, cursor - 1)
            self._recompute(wb_id)
        ops = self._ops(wb_id)
        return f"已撤销：{ops[cursor - 1]['description']}"

    def redo(self, wb_id: str) -> str:
        with WorkbookManager._write_lock:
            cursor = self.storage.get_cursor(wb_id)
            ops = self._ops(wb_id)
            if cursor >= len(ops):
                raise mutations.OpError("没有可重做的操作")
            self.storage.set_cursor(wb_id, cursor + 1)
            self._recompute(wb_id)
        return f"已重做：{ops[cursor]['description']}"

    def goto(self, wb_id: str, n: int) -> str:
        """回到第 n 步操作之后的状态（n=0 即初始导入态）。"""
        with WorkbookManager._write_lock:
            ops = self._ops(wb_id)
            if n < 0 or n > len(ops):
                raise mutations.OpError(f"步骤 {n} 超出范围（当前共 {len(ops)} 步）")
            self.storage.set_cursor(wb_id, n)
            self._recompute(wb_id)
        if n == 0:
            return "已回到初始状态（v0，导入原始数据）"
        return f"已回到第 {n} 步之后：{ops[n - 1]['description']}"

    def save_snapshot(self, wb_id: str, name: str) -> str:
        version_map = {s["sheet_name"]: int(s["current_version"])
                       for s in self.storage.list_sheets(wb_id, only_active=False)}
        self.storage.add_snapshot(wb_id, name, version_map)
        return f"快照「{name}」已保存"

    def restore_snapshot(self, wb_id: str, name: str) -> str:
        snap = self.storage.get_snapshot(wb_id, name)
        if snap is None:
            raise mutations.OpError(f"快照「{name}」不存在")
        version_map = _loads(snap["version_map_json"])
        with WorkbookManager._write_lock:
            # 找到与该快照匹配的游标：逐候选游标重放对比（操作数少，代价可忽略）
            target_cursor = None
            saved_ops = self._ops(wb_id)
            for cand in range(len(saved_ops), -1, -1):
                self.storage.set_cursor(wb_id, cand)
                self._recompute(wb_id)
                cur_map = {s["sheet_name"]: int(s["current_version"])
                           for s in self.storage.list_sheets(wb_id, only_active=False)}
                if cur_map == {k: v for k, v in version_map.items() if k in cur_map}:
                    target_cursor = cand
                    break
            if target_cursor is None:
                raise mutations.OpError(f"快照「{name}」对应的版本链状态已不存在（可能链被新操作截断）")
        return f"已恢复到快照「{name}」"

    # ------------------------------------------------------------------
    # 样式
    # ------------------------------------------------------------------
    def get_styles(self, wb_id: str, sheet: str) -> dict[str, dict]:
        return self.storage.get_styles(wb_id, sheet)

    # ------------------------------------------------------------------
    # 导出
    # ------------------------------------------------------------------
    def export_xlsx(self, wb_id: str, out_path: str | Path | "io.BytesIO") -> Path | "io.BytesIO":
        """导出当前状态为带样式 xlsx（支持传 BytesIO 直接拿字节）。

        AC-5 保底：所有 sheet 都在 v0 且 workbook 源自上传/示例时，
        直接复制原始文件（与用户上传字节一致）。
        """
        import io as _io
        if not isinstance(out_path, _io.BytesIO):
            out_path = Path(out_path)
            out_path.parent.mkdir(parents=True, exist_ok=True)
        wb = self.storage.get_workbook(wb_id)
        if wb is None:
            raise WorkbookNotFoundError(f"workbook {wb_id} 不存在")
        sheets = self.storage.list_sheets(wb_id, only_active=False)
        all_v0 = all(int(s["current_version"]) == 0 for s in sheets)
        if all_v0 and wb["original_path"] and Path(wb["original_path"]).exists():
            if isinstance(out_path, _io.BytesIO):
                out_path.write(Path(wb["original_path"]).read_bytes())
            else:
                shutil.copyfile(wb["original_path"], out_path)
            return out_path

        from openpyxl import Workbook as XlsxWorkbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        xw = XlsxWorkbook()
        xw.remove(xw.active)
        for s in self.storage.list_sheets(wb_id, only_active=True):
            name = s["sheet_name"]
            df = self.get_df(wb_id, name)
            ws = xw.create_sheet(title=name[:31])
            ws.append([str(c) for c in df.columns])
            for c_idx in range(1, len(df.columns) + 1):
                cell = ws.cell(row=1, column=c_idx)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="DDEBF7")
            for row in df.itertuples(index=False):
                ws.append([_excel_value(v) for v in row])
            styles = self.get_styles(wb_id, name)
            for ref, style in styles.items():
                _apply_style(ws, df, ref, style, Font, PatternFill, Alignment, get_column_letter)
            for i, col in enumerate(df.columns, start=1):
                col_styles = [st for r, st in styles.items() if r == f"col:{col}"]
                if col_styles and col_styles[-1].get("width"):
                    ws.column_dimensions[get_column_letter(i)].width = float(col_styles[-1]["width"])
                elif len(df):
                    width = max(10, min(40, int(df[col].astype(str).str.len().max() * 1.6) if len(df) else 10))
                    ws.column_dimensions[get_column_letter(i)].width = width
        xw.save(out_path)
        return out_path

    def export_xlsx_bytes(self, wb_id: str) -> bytes:
        import io
        buf = io.BytesIO()
        self.export_xlsx(wb_id, buf)
        return buf.getvalue()


def _excel_value(v):
    import math
    if v is None:
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    try:
        pd_isna = pd.isna(v)
        if pd_isna is True:
            return None
    except (TypeError, ValueError):
        pass
    return v


def _apply_style(ws, df: pd.DataFrame, ref: str, style: dict,
                 Font, PatternFill, Alignment, get_column_letter) -> None:
    """把一条样式记录（'B2' / 'A1:C10' / 'col:列名'）应用到 openpyxl 工作表。"""
    import re
    try:
        if ref.startswith("col:"):
            col_name = ref[4:]
            if col_name in list(df.columns):
                c_idx = list(df.columns).index(col_name) + 1
                if style.get("width"):
                    ws.column_dimensions[get_column_letter(c_idx)].width = float(style["width"])
                rng = ws.iter_cols(min_col=c_idx, max_col=c_idx, min_row=1, max_row=len(df) + 1)
                for col_cells in rng:
                    for cell in col_cells:
                        _style_cell(cell, style, Font, PatternFill, Alignment)
            return
        m = re.match(r"^([A-Z]+\d+)(?::([A-Z]+\d+))?$", ref)
        if not m:
            return
        start, end = m.group(1), m.group(2) or m.group(1)
        for row in ws[start:end]:
            for cell in row:
                _style_cell(cell, style, Font, PatternFill, Alignment)
    except Exception:
        # 样式应用失败不阻断导出（数据完整性优先）
        pass


def _style_cell(cell, style: dict, Font, PatternFill, Alignment) -> None:
    if style.get("bold") is not None or style.get("font_color"):
        cell.font = Font(
            bold=bool(style.get("bold")),
            color=style.get("font_color") or "000000",
        )
    if style.get("bg_color"):
        cell.fill = PatternFill("solid", fgColor=str(style["bg_color"]).lstrip("#").upper())
    if style.get("number_format"):
        cell.number_format = style["number_format"]
    if style.get("align"):
        cell.alignment = Alignment(horizontal=style["align"])


def _dumps(obj) -> str:
    import json
    return json.dumps(obj, ensure_ascii=False)


def _loads(s: str):
    import json
    return json.loads(s)
